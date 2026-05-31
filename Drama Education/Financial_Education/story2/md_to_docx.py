import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def parse_markdown(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    return lines

def process_bold(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)

def create_docx(source_file, target_file):
    doc = Document()
    lines = parse_markdown(source_file)
    
    in_table = False
    table_lines = []
    
    for line in lines:
        line = line.strip()
        
        # Handle Table
        if line.startswith('|'):
            in_table = True
            table_lines.append(line)
            continue
        else:
            if in_table:
                # Process the collected table
                process_table(doc, table_lines)
                table_lines = []
                in_table = False
        
        if not line:
            continue
            
        # Headers
        if line.startswith('#'):
            level = len(line.split(' ')[0])
            text = line.lstrip('#').strip()
            # doc.add_heading(text, level=level) # Heading level is limited 0-9
            # Customizing heading to look better or just use default
            if level <= 9:
                h = doc.add_heading(text, level=level)
            else:
                p = doc.add_paragraph(text)
                p.bold = True
        
        # Blockquotes
        elif line.startswith('> '):
            text = line[2:]
            p = doc.add_paragraph()
            process_bold(p, text)
            p.style = 'Intense Quote'
            
        # List items
        elif line.startswith('* ') or line.startswith('- '):
            text = line[2:]
            p = doc.add_paragraph(style='List Bullet')
            process_bold(p, text)
            
        # Normal text (handling bold)
        else:
            p = doc.add_paragraph()
            process_bold(p, line)
            
    # If file ends with table
    if in_table:
        process_table(doc, table_lines)
        
    doc.save(target_file)
    print(f"File saved to {target_file}")

def process_table(doc, table_lines):
    # Filter out separator lines like |---|---|
    data_lines = [l for l in table_lines if '---' not in l]
    if not data_lines:
        return
        
    # Get columns count from first line
    header_cells = [c.strip() for c in data_lines[0].split('|') if c]
    cols_count = len(header_cells)
    
    table = doc.add_table(rows=len(data_lines), cols=cols_count)
    table.style = 'Table Grid'
    
    for r_idx, row_line in enumerate(data_lines):
        cells = [c.strip() for c in row_line.split('|')][1:-1] # Skip empty start/end logic if | is at ends
        # Adjustment for split logic depending on if line ends with |
        if row_line.endswith('|'):
            # split results in ['', 'c1', 'c2', '']
            cells = [c.strip() for c in row_line.split('|') if c.strip() or c == ' '][0:cols_count]
        else:
             cells = [c.strip() for c in row_line.split('|') if c.strip()][0:cols_count]
             
        row_cells = table.rows[r_idx].cells
        for c_idx, cell_text in enumerate(cells):
            if c_idx < len(row_cells):
                # Handle <br> in table cells
                parts = cell_text.split('<br>')
                p = row_cells[c_idx].paragraphs[0]
                for i, part in enumerate(parts):
                    process_bold(p, part.strip())
                    if i < len(parts) - 1:
                        p.add_run('\n')

source = "lesson_plan_activity2_rainbow_port.md"
target = "lesson_plan_activity2_rainbow_port_v2.docx"
create_docx(source, target)
